from __future__ import annotations

from datetime import date
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "apps" / "talisma_sis" / "docs" / "client" / "Bryan_University_Project_Handbook.docx"
SNAPSHOT_DATE = "14 July 2026"
BRANCH = "feat/demo-mvp"
COMMIT = "c1f880f"

NAVY = RGBColor(16, 43, 78)
BLUE = RGBColor(46, 116, 181)
GOLD = RGBColor(217, 164, 65)
INK = RGBColor(36, 54, 75)
GRAY = RGBColor(86, 96, 105)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = "E8EEF5"
LIGHT_GOLD = "FFF8E8"
LIGHT_RED = "FCE8E6"
LIGHT_GRAY = "F2F4F7"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_font(run, size=11, bold=False, color=INK, italic=False, name="Calibri"):
	run.font.name = name
	run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
	run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
	run.font.size = Pt(size)
	run.bold = bold
	run.italic = italic
	if color:
		run.font.color.rgb = color
	return run


def shade(cell, fill):
	tc_pr = cell._tc.get_or_add_tcPr()
	shd = tc_pr.find(qn("w:shd"))
	if shd is None:
		shd = OxmlElement("w:shd")
		tc_pr.append(shd)
	shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
	tc_pr = cell._tc.get_or_add_tcPr()
	tc_w = tc_pr.find(qn("w:tcW"))
	if tc_w is None:
		tc_w = OxmlElement("w:tcW")
		tc_pr.append(tc_w)
	tc_w.set(qn("w:w"), str(width))
	tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
	tc = cell._tc
	tc_pr = tc.get_or_add_tcPr()
	tc_mar = tc_pr.first_child_found_in("w:tcMar")
	if tc_mar is None:
		tc_mar = OxmlElement("w:tcMar")
		tc_pr.append(tc_mar)
	for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
		node = tc_mar.find(qn(f"w:{tag}"))
		if node is None:
			node = OxmlElement(f"w:{tag}")
			tc_mar.append(node)
		node.set(qn("w:w"), str(value))
		node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT):
	assert sum(widths) == TABLE_WIDTH
	table.autofit = False
	table.alignment = WD_TABLE_ALIGNMENT.LEFT
	tbl_pr = table._tbl.tblPr
	tbl_w = tbl_pr.find(qn("w:tblW"))
	if tbl_w is None:
		tbl_w = OxmlElement("w:tblW")
		tbl_pr.append(tbl_w)
	tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
	tbl_w.set(qn("w:type"), "dxa")
	tbl_ind = tbl_pr.find(qn("w:tblInd"))
	if tbl_ind is None:
		tbl_ind = OxmlElement("w:tblInd")
		tbl_pr.append(tbl_ind)
	tbl_ind.set(qn("w:w"), str(indent))
	tbl_ind.set(qn("w:type"), "dxa")
	grid = table._tbl.tblGrid
	for child in list(grid):
		grid.remove(child)
	for width in widths:
		col = OxmlElement("w:gridCol")
		col.set(qn("w:w"), str(width))
		grid.append(col)
	for row in table.rows:
		for idx, cell in enumerate(row.cells):
			set_cell_width(cell, widths[idx])
			set_cell_margins(cell)
			cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_table_header(row):
	tr_pr = row._tr.get_or_add_trPr()
	tbl_header = OxmlElement("w:tblHeader")
	tbl_header.set(qn("w:val"), "true")
	tr_pr.append(tbl_header)


def prevent_row_split(row):
	tr_pr = row._tr.get_or_add_trPr()
	cant_split = OxmlElement("w:cantSplit")
	tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths, font_size=9.2):
	table = doc.add_table(rows=1, cols=len(headers))
	table.style = "Table Grid"
	for idx, header in enumerate(headers):
		cell = table.rows[0].cells[idx]
		shade(cell, LIGHT_BLUE)
		p = cell.paragraphs[0]
		p.paragraph_format.space_after = Pt(1)
		set_font(p.add_run(str(header)), size=9.5, bold=True, color=NAVY)
	repeat_table_header(table.rows[0])
	for values in rows:
		row = table.add_row()
		prevent_row_split(row)
		for idx, value in enumerate(values):
			cell = row.cells[idx]
			p = cell.paragraphs[0]
			p.paragraph_format.space_after = Pt(1)
			p.paragraph_format.line_spacing = 1.08
			set_font(p.add_run(str(value)), size=font_size)
	set_table_geometry(table, widths)
	spacer = doc.add_paragraph()
	spacer.paragraph_format.space_after = Pt(2)
	return table


def add_callout(doc, title, text, fill=LIGHT_GOLD):
	p = doc.add_paragraph()
	p_pr = p._p.get_or_add_pPr()
	shd = OxmlElement("w:shd")
	shd.set(qn("w:fill"), str(fill))
	p_pr.append(shd)
	p_bdr = OxmlElement("w:pBdr")
	left = OxmlElement("w:left")
	left.set(qn("w:val"), "single")
	left.set(qn("w:sz"), "18")
	left.set(qn("w:space"), "8")
	left.set(qn("w:color"), str(GOLD))
	p_bdr.append(left)
	p_pr.append(p_bdr)
	p.paragraph_format.left_indent = Pt(10)
	p.paragraph_format.right_indent = Pt(8)
	p.paragraph_format.space_before = Pt(4)
	p.paragraph_format.space_after = Pt(5)
	p.paragraph_format.line_spacing = 1.15
	set_font(p.add_run(f"{title}: "), bold=True, color=NAVY)
	set_font(p.add_run(text))


def add_p(doc, text="", bold_prefix=None, italic=False, after=6):
	p = doc.add_paragraph()
	p.paragraph_format.space_after = Pt(after)
	p.paragraph_format.line_spacing = 1.25
	if bold_prefix and text.startswith(bold_prefix):
		set_font(p.add_run(bold_prefix), bold=True, color=NAVY)
		set_font(p.add_run(text[len(bold_prefix):]), italic=italic)
	else:
		set_font(p.add_run(text), italic=italic)
	return p


def add_bullets(doc, items, level=0):
	for item in items:
		p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
		p.paragraph_format.space_after = Pt(4)
		p.paragraph_format.line_spacing = 1.25
		set_font(p.add_run(str(item)))


def add_steps(doc, items):
	for item in items:
		p = doc.add_paragraph(style="List Number")
		p.paragraph_format.space_after = Pt(4)
		p.paragraph_format.line_spacing = 1.25
		set_font(p.add_run(str(item)))


def add_heading(doc, text, level=1):
	p = doc.add_paragraph(text, style=f"Heading {level}")
	p.paragraph_format.keep_with_next = True
	return p


def add_code(doc, lines):
	for line in lines.strip("\n").splitlines():
		p = doc.add_paragraph(style="Code Block")
		set_font(p.add_run(line or " "), size=8.5, color=INK, name="Consolas")


def page_break(doc):
	doc.add_page_break()


def add_page_field(paragraph):
	run = paragraph.add_run()
	fld_char = OxmlElement("w:fldChar")
	fld_char.set(qn("w:fldCharType"), "begin")
	instr = OxmlElement("w:instrText")
	instr.set(qn("xml:space"), "preserve")
	instr.text = " PAGE "
	sep = OxmlElement("w:fldChar")
	sep.set(qn("w:fldCharType"), "separate")
	text = OxmlElement("w:t")
	text.text = "1"
	end = OxmlElement("w:fldChar")
	end.set(qn("w:fldCharType"), "end")
	for node in (fld_char, instr, sep, text, end):
		run._r.append(node)
	set_font(run, size=8.5, color=GRAY)


def configure_styles(doc):
	styles = doc.styles
	normal = styles["Normal"]
	normal.font.name = "Calibri"
	normal.font.size = Pt(11)
	normal.font.color.rgb = INK
	normal.paragraph_format.space_after = Pt(6)
	normal.paragraph_format.line_spacing = 1.25
	for level, size, before, after, color in (
		(1, 16, 18, 10, BLUE),
		(2, 13, 14, 7, BLUE),
		(3, 12, 10, 5, NAVY),
	):
		style = styles[f"Heading {level}"]
		style.font.name = "Calibri"
		style.font.size = Pt(size)
		style.font.bold = True
		style.font.color.rgb = color
		style.paragraph_format.space_before = Pt(before)
		style.paragraph_format.space_after = Pt(after)
		style.paragraph_format.keep_with_next = True
	for name, left, first in (("List Bullet", 540, -270), ("List Bullet 2", 900, -270), ("List Number", 540, -270)):
		style = styles[name]
		style.font.name = "Calibri"
		style.font.size = Pt(11)
		style.paragraph_format.left_indent = Pt(left / 20)
		style.paragraph_format.first_line_indent = Pt(first / 20)
		style.paragraph_format.space_after = Pt(4)
		style.paragraph_format.line_spacing = 1.25
	code = styles.add_style("Code Block", 1)
	code.font.name = "Consolas"
	code.font.size = Pt(8.5)
	code.paragraph_format.left_indent = Inches(0.18)
	code.paragraph_format.right_indent = Inches(0.08)
	code.paragraph_format.space_before = Pt(0)
	code.paragraph_format.space_after = Pt(0)
	code.paragraph_format.line_spacing = 1.05
	p_pr = code.element.get_or_add_pPr()
	shd = OxmlElement("w:shd")
	shd.set(qn("w:fill"), LIGHT_GRAY)
	p_pr.append(shd)


def add_running_furniture(doc):
	section = doc.sections[0]
	section.top_margin = Inches(1)
	section.bottom_margin = Inches(1)
	section.left_margin = Inches(1)
	section.right_margin = Inches(1)
	section.header_distance = Inches(0.492)
	section.footer_distance = Inches(0.492)
	section.different_first_page_header_footer = True
	header = section.header.paragraphs[0]
	header.alignment = WD_ALIGN_PARAGRAPH.LEFT
	set_font(header.add_run("BRYAN UNIVERSITY  |  PROJECT HANDBOOK"), size=8.5, bold=True, color=GRAY)
	footer = section.footer.paragraphs[0]
	footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
	set_font(footer.add_run(f"Working-tree snapshot | {SNAPSHOT_DATE}  |  Page "), size=8.5, color=GRAY)
	add_page_field(footer)
	first_footer = section.first_page_footer.paragraphs[0]
	first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
	set_font(first_footer.add_run("Internal project reference | Synthetic demo data only"), size=8.5, color=GRAY)


def add_cover(doc):
	p = doc.add_paragraph()
	p.paragraph_format.space_before = Pt(88)
	p.paragraph_format.space_after = Pt(18)
	p.alignment = WD_ALIGN_PARAGRAPH.CENTER
	set_font(p.add_run("PROJECT HANDBOOK"), size=10, bold=True, color=GOLD)
	title = doc.add_paragraph()
	title.alignment = WD_ALIGN_PARAGRAPH.CENTER
	title.paragraph_format.space_after = Pt(8)
	set_font(title.add_run("Bryan University SIS"), size=30, bold=True, color=NAVY)
	subtitle = doc.add_paragraph()
	subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
	subtitle.paragraph_format.space_after = Pt(28)
	set_font(subtitle.add_run("Architecture, configuration, workflows, operations, data model, and delivery reference"), size=14, color=GRAY)
	add_table(doc, ["Snapshot", "Branch / commit", "Coverage"], [[SNAPSHOT_DATE, f"{BRANCH} / {COMMIT}", "Current working tree and live local demo"]], [2200, 3000, 4160])
	add_callout(doc, "Purpose", "A single comprehensive handover for product owners, architects, developers, administrators, implementation teams, testers, and client-demo operators.")
	add_callout(doc, "Security", "Credentials, passwords, tokens, and secret values are intentionally not reproduced. Their configuration locations and rotation requirements are documented instead.", LIGHT_RED)


def add_contents(doc):
	page_break(doc)
	add_heading(doc, "Contents", 1)
	sections = [
		"1. Executive overview", "2. Current implementation status and scope", "3. Architecture and ownership model",
		"4. Runtime topology and technology stack", "5. Repository organization", "6. Application extension points and hooks",
		"7. User experience, branding, and navigation", "8. Data model: Talisma-owned DocTypes",
		"9. Data model: extended Education records", "10. Admissions workflow", "11. Registrar and section-registration workflow",
		"12. US academic logic", "13. Assessment, grading, and report cards", "14. Institutional hierarchy and closure materialization",
		"15. Demo dataset and workspaces", "16. Configuration reference", "17. Security, privacy, and permissions",
		"18. Build, deployment, and site operations", "19. Testing, health checks, migrations, and recovery",
		"20. Troubleshooting and support", "21. Client demonstration runbook", "22. Production-readiness gaps and roadmap",
		"Appendices: APIs, file map, glossary, and complete manifest",
	]
	add_bullets(doc, sections)
	add_callout(doc, "Reading guide", "Sections 1-4 orient decision-makers; 8-14 describe business and system logic; 16-20 are the operator reference; appendices are intended for engineering handover.", LIGHT_BLUE)


def add_main_content(doc):
	page_break(doc)
	add_heading(doc, "1. Executive overview", 1)
	add_p(doc, "Talisma SIS is an upgrade-safe higher-education Student Information System implemented as a first-class Frappe application alongside ERPNext and Frappe Education. The current client-facing experience is branded as Bryan University. The product direction is an institutional system of record spanning admission through graduation, with ERPNext retained as the authority for accounting and enterprise records.")
	add_table(doc, ["Dimension", "Current position"], [
		["Product", "Higher-education SIS foundation plus a synthetic US client-demo vertical slice"],
		["App", "talisma_sis 0.0.1"], ["Platform", "Frappe 16.25.0, ERPNext 16.26.2, Education 16.0.1"],
		["Demo identity", "Bryan University workspace; demo.talisma.local; localhost:8082"],
		["Persistence", "MariaDB 11.8, Frappe DocTypes, private/public files"], ["Async / realtime", "Redis 8.6, workers, scheduler, websocket"],
		["Design rule", "No modifications to Frappe, ERPNext, or Education core"],
	], [2200, 7160])
	add_callout(doc, "Most important caveat", "The institutional hierarchy foundation has controller-level invariants and tests. Admissions, registration, US academic extensions, branding, and report-card behavior are intentionally restricted to the isolated demo site and must not be represented as production-complete.", LIGHT_RED)

	add_heading(doc, "2. Current implementation status and scope", 1)
	add_heading(doc, "2.1 Implemented Talisma foundation", 2)
	add_bullets(doc, [
		"Institution, Campus, Academic Unit Type, Academic Unit, Structure Version, Unit Placement, Closure Build, and Unit Closure models.",
		"Effective-date validation, immutable identifiers, scoped uniqueness, lifecycle rules, submitted-structure governance, forest validation, and closure-table integrity.",
		"Manager, Viewer, Academic Structure Approver, and System Manager permission boundaries.",
		"Version-controlled schema constraints and post-model-sync patches.",
	])
	add_heading(doc, "2.2 Implemented isolated demo slice", 2)
	add_bullets(doc, [
		"Bryan University branding, eight module launchers, curated workspaces, sidebars, quick actions, high-contrast forms and lists, and responsive full-page form treatment.",
		"Admissions conversion from approved applicant to Student and Program Enrollment with duplicate prevention and linked decision data.",
		"Program and section registration through standard Course Enrollment and Student Group roster records, including idempotence and capacity checks.",
		"US catalog, term-milestone, CRN, section, meeting, grading-range, and student-profile extensions.",
		"Education report-card compatibility adapter for Frappe v16-safe attendance aggregation and isolated PDF generation.",
	])
	add_heading(doc, "2.3 Explicitly deferred", 2)
	add_bullets(doc, [
		"Production FERPA consent, directory-information choices, proxy access, legitimate-interest audit, and institutional scope enforcement.",
		"Federal Student Aid/FAFSA/ISIR/COD/NSLDS/SAP, SEVIS, 1098-T, IPEDS, Clearinghouse, Title IX, ADA, NCAA, and state authorization processing.",
		"Official GPA, repeat and transfer policy, curriculum versioning, degree audit, transcript control, conferral, and production billing residency/tuition calculation.",
		"Production portal, integration contracts, high-volume reporting store, disaster recovery certification, and complete accessibility validation.",
	])
	add_callout(doc, "Source-of-truth note", "Several architecture documents predate the demo increment and correctly describe features as gated or proposed. This handbook reports both the approved foundation and the newer working-tree demo behavior.", LIGHT_BLUE)

	page_break(doc)
	add_heading(doc, "3. Architecture and ownership model", 1)
	add_table(doc, ["Layer", "Primary responsibility", "Owned by"], [
		["Experience", "Desk, workspaces, forms, lists, portals, reports, print output", "Talisma plus Frappe UI"],
		["Application", "Services, validations, orchestration, whitelisted operations", "talisma_sis"],
		["Domain", "Admissions, student lifecycle, academic policy, structure", "Talisma / reviewed Education reuse"],
		["Platform", "ORM, permissions, audit, jobs, cache, realtime, metadata", "Frappe"],
		["Enterprise", "Company, workforce, receivables, payments, ledger", "ERPNext"],
		["Academic upstream", "Student, Program, Course, term, attendance, assessment", "Frappe Education, extended by Talisma"],
		["Data", "MariaDB records, site files, Redis queues/cache", "Frappe runtime"],
	], [1500, 4860, 3000])
	add_heading(doc, "3.1 Extension priority", 2)
	add_steps(doc, [
		"Configure an upstream capability when its lifecycle and ownership match.",
		"Extend a standard DocType with supported hooks or app-owned custom fields.",
		"Create a Talisma DocType when the concept needs distinct identity, lifecycle, permissions, or retention.",
		"Add an application service when a transaction spans multiple records.",
		"Override upstream behavior only with a documented compatibility reason and regression coverage.",
	])
	add_heading(doc, "3.2 Authoritative ownership", 2)
	add_table(doc, ["Concern", "Authority"], [
		["Authentication and authorization", "Frappe User, Role, User Permission, DocPerm, sessions"],
		["Legal/accounting entity and ledger", "ERPNext Company and Accounts"],
		["Academic identity", "Education Student/Instructor linked through approved party model"],
		["Institutional academic organization", "Talisma Institution, Campus, Academic Unit, Structure Version"],
		["Demo programs/courses/terms", "Education records with Talisma-owned extension fields"],
		["Final higher-education policy", "Future Talisma services and effective-dated records"],
	], [3000, 6360])

	add_heading(doc, "4. Runtime topology and technology stack", 1)
	add_table(doc, ["Service", "Role", "Dependency / port"], [
		["frontend", "Nginx reverse proxy and static assets", "127.0.0.1:8082 -> 8080"],
		["backend", "Gunicorn/Frappe web application", "Internal 8000; health /api/method/ping"],
		["websocket", "Frappe realtime Socket.IO", "Internal 9000"],
		["queue-short", "Short/default jobs", "Redis queue"], ["queue-long", "Long/default/short jobs", "Redis queue"],
		["scheduler", "Scheduled task dispatcher", "Frappe scheduler"], ["configurator", "Writes shared DB/Redis/socket settings", "Runs before application services"],
		["db", "MariaDB transactional storage", "MariaDB 11.8, isolated named volume"],
		["redis-cache", "Application cache", "Redis 8.6 Alpine"], ["redis-queue", "Queues and durable queue data", "Redis 8.6 Alpine, named volume"],
	], [1700, 4300, 3360])
	add_callout(doc, "Isolation", "Compose project talisma-demo creates its own network and named volumes. The demo frontend binds only to loopback. Synthetic data only; never seed talisma.local.")

	add_heading(doc, "5. Repository organization", 1)
	add_table(doc, ["Path", "Purpose"], [
		["apps/talisma_sis/talisma_sis", "Application Python, hooks, services, metadata, assets, translations"],
		["apps/talisma_sis/talisma_sis/talisma_sis/doctype", "Talisma-owned DocType schemas, controllers, tests"],
		["apps/talisma_sis/demo", "Isolated demo configuration, seed/runbook, scope, screenshots"],
		["apps/talisma_sis/docs", "Vision, architecture, ADRs, module specifications, client guides"],
		["apps/talisma_sis/tests/compatibility/education", "Pinned Education evaluation image and Compose overlay"],
		["apps/education", "Pinned upstream Education source included in this working tree"],
		["images", "Frappe/ERPNext layered and Talisma overlay image recipes"],
		["compose.yaml + overrides", "Base service topology and database/Redis/deployment variants"],
		["docs", "frappe_docker operations and deployment documentation"],
	], [3900, 5460])

	page_break(doc)
	add_heading(doc, "6. Application extension points and hooks", 1)
	add_table(doc, ["Hook", "Configured behavior"], [
		["app_include_css / web_include_css", "Loads talisma_demo.css"],
		["app_include_js", "Loads talisma_demo.js"],
		["doctype_js", "Student Applicant, Student, Program Enrollment, and report-card tool client behavior"],
		["doc_events", "Server validation for Course, Academic Term, and Student Group"],
		["override_whitelisted_methods", "Routes Education report-card preview to the compatibility adapter"],
		["after_request", "Redirects authenticated demo Desk root to Bryan University"],
	], [3100, 6260])
	add_p(doc, "All demo-only server services call a site guard and fail closed unless frappe.local.site equals demo.talisma.local. Client scripts also check frappe.boot.sitename before enabling workflow buttons or demo queries.")

	add_heading(doc, "7. User experience, branding, and navigation", 1)
	add_heading(doc, "7.1 Branding", 2)
	add_bullets(doc, [
		"Workspace and browser identity: Bryan University.", "Primary palette: navy #102B4E, gold #FFC629, ink #24364B.",
		"Global university mark, title prefix, favicon, splash image, module icons, and default workspace are configured by the demo setup.",
		"Legacy Talisma University and misspelled Bryne University workspace names are migrated where found.",
	])
	add_heading(doc, "7.2 Module launchers", 2)
	add_table(doc, ["Module", "Primary purpose"], [
		["Admissions", "Admission Intake and Student Application"], ["Student Records", "Students and enrollment records"],
		["Registrar", "Program/course registration and registrar number cards"], ["Academics", "Academic year, term, program, course, topic"],
		["Faculty & Sections", "Instructor, course sections, schedules, rooms"], ["Assessment & Grades", "Plans, results, scales, report cards"],
		["Attendance", "Attendance, tool, leave applications"], ["Reports", "Assessment, grades, attendance, and contact reports"],
	], [2700, 6660])
	add_heading(doc, "7.3 Visual-system rules", 2)
	add_bullets(doc, [
		"Navy sidebar with white text and gold active item; high-contrast labels, headers, controls, checkboxes, and dividers.",
		"List headers remain navy; yellow is used as an accent and not as a data-state color.",
		"Forms use responsive full-page cards; short forms fill the viewport while long forms grow naturally.",
		"Admission Intake uses a wide content layout; its unused Introduction editor is hidden.",
		"Dialogs intentionally remove excess nested borders; editable child grids retain compact native table styling.",
		"Toolbar filters, sort controls, text inputs, and operator controls receive visible borders without altering business semantics.",
	])

	add_heading(doc, "8. Data model: Talisma-owned DocTypes", 1)
	add_table(doc, ["DocType", "Purpose", "Lifecycle / key rule"], [
		["Talisma Institution", "Institution identity and locale", "UUID; immutable normalized code; effective dates"],
		["Talisma Campus", "Physical/virtual delivery location", "Institution-scoped code; address for non-virtual; cannot reopen Closed"],
		["Talisma Academic Unit Type", "Governed unit classification and capability ceiling", "Globally unique code; referenced configuration becomes immutable"],
		["Talisma Academic Unit", "College/school/department identity", "Institution-scoped code; capability cannot exceed type; effective lifecycle"],
		["Talisma Structure Version", "Submitted effective-dated academic hierarchy", "Forest validation; bounded interval; no cancel; successor version"],
		["Talisma Unit Placement", "Places one unit under another in a structure", "Draft-only mutation; same institution; immutable links"],
		["Talisma Closure Build", "Materialization run and evidence", "Service-managed Building/Ready/Failed/Superseded"],
		["Talisma Unit Closure", "Ancestor/descendant/depth materialized rows", "Read-only service-managed closure table"],
		["Talisma Family Member", "Embedded family/relationship row", "Child table reused by applicant and student"],
	], [2400, 3500, 3460], font_size=8.8)
	add_heading(doc, "8.1 Naming and integrity", 2)
	add_bullets(doc, [
		"Primary records use UUID4 names; business codes are normalized and validated separately.",
		"Database patches add composite unique constraints for campus, academic unit, structure, closure, and supersedes relationships.",
		"track_changes is enabled on governed master and structure records.",
		"Effective intervals must fit within the owning institution and, where applicable, the referenced unit.",
	])
	add_heading(doc, "8.2 Roles", 2)
	add_table(doc, ["Role", "Authority"], [
		["Talisma Institution Manager", "Create and maintain institution, campus, academic-unit and draft structure records"],
		["Talisma Institution Viewer", "Read-only access to institutional foundation"],
		["Talisma Academic Structure Approver", "Submit governed Structure Versions; cannot freely alter supporting masters"],
		["System Manager", "Administrative control and service/support visibility"],
	], [3200, 6160])

	page_break(doc)
	add_heading(doc, "9. Data model: extended Education records", 1)
	add_table(doc, ["Standard DocType", "Talisma-owned fields / behavior"], [
		["Student Applicant", "Campus; decision date/note; converted Student; Program Enrollment; admit type; residency; deposit; family members"],
		["Student", "Home campus; preferred name; pronouns; family; academic level; class standing; enrollment status; catalog year; expected graduation term; advisor"],
		["Program", "Owning academic unit"],
		["Course", "Academic unit; subject code; catalog number; credit hours; level; grading basis; effective term; repeatable flag"],
		["Academic Term", "Registration opens; add/drop; census; withdrawal; grades due"],
		["Student Group", "CRN; section; campus; delivery; status; waitlist; meeting days/times; room; primary instructor"],
		["Course Enrollment", "Course section / CRN link"], ["Instructor", "Academic unit"],
		["Program Enrollment", "Campus; calculated registration status; registration-updated timestamp"],
	], [2500, 6860], font_size=8.8)
	add_heading(doc, "9.1 Metadata/property changes", 2)
	add_bullets(doc, [
		"Student and Student Applicant family sections are relabeled Family Details and legacy Guardian tables are hidden in the demo experience.",
		"Student Admission is presented as Admission Intake; its Introduction editor is hidden.",
		"Grading Scale Interval Threshold is relabeled Minimum Score, and Grade Description is relabeled Score Range.",
		"Demo grading rows remain numeric minimum thresholds for calculation and display explicit ranges in the adjacent description field.",
	])

	add_heading(doc, "10. Admissions workflow", 1)
	add_callout(doc, "Flow", "Approved Student Applicant -> Admit and Create Student -> Student -> Program Enrollment -> applicant links and decision audit", LIGHT_BLUE)
	add_steps(doc, [
		"Admissions user opens an approved Student Applicant.", "Client script displays Admit and Create Student only when status is Approved and no converted student is linked.",
		"User optionally enters an internal decision note and confirms.", "Server verifies demo site, write permission, and Approved/Admitted state.",
		"Service reuses a linked Student, an applicant-linked Student, or an email-matched Student; otherwise it maps a new Student and copies campus and family rows.",
		"Service reuses or creates Program Enrollment for student, program, year, and optional term.",
		"Applicant becomes Admitted and stores decision date, note, Student link, and Program Enrollment link.",
		"UI reloads the applicant and routes to the Program Enrollment.",
	])
	add_heading(doc, "10.1 Idempotence and failure rules", 2)
	add_bullets(doc, [
		"Repeated admission returns the same linked Student and Program Enrollment instead of duplicating records.",
		"Applied or Rejected applicants cannot be admitted until status is Approved.",
		"Server permission is authoritative; the presence of a button is not authorization.",
		"The conversion is a demo workflow and is not yet a production identity-resolution or offer-acceptance engine.",
	])

	add_heading(doc, "11. Registrar and section-registration workflow", 1)
	add_callout(doc, "Flow", "Program Enrollment -> allowed Program Courses -> open term Sections -> Course Enrollment + Student Group roster -> registration summary", LIGHT_BLUE)
	add_steps(doc, [
		"Open a saved Program Enrollment and choose Register Courses.", "Client requests open section options for the same program, academic year, and term.",
		"Server returns only Course-based, enabled, Open Student Groups whose Course is in the Program curriculum.",
		"User selects sections by CRN and meeting context.", "Server validates permission, non-cancelled enrollment, curriculum membership, term match, open status, and section capacity.",
		"Service creates or links one Course Enrollment per course and associates the section.", "Student is added once to the Student Group roster.",
		"Program Enrollment receives a calculated registration label and update timestamp.",
	])
	add_table(doc, ["Rule", "Current demo behavior"], [
		["Maximum load", "Five courses per term in register_courses; section registration relies on the prepared workflow and available curriculum"],
		["Duplicate course", "No duplicate Course Enrollment; a different section of the same course is rejected"],
		["Capacity", "Active roster count must be below max_strength"], ["Waitlist", "Capacity metadata exists, but automatic waitlist processing is deferred"],
		["Holds/prerequisites/conflicts", "Not yet implemented as production registration policy"],
	], [2500, 6860])

	page_break(doc)
	add_heading(doc, "12. US academic logic", 1)
	add_heading(doc, "12.1 Course validation", 2)
	add_bullets(doc, ["Course Code, Course Name, and Course Type are required.", "Credit Hours must be greater than zero."])
	add_heading(doc, "12.2 Academic-term chronology", 2)
	add_p(doc, "All milestones are required and must be chronological:")
	add_steps(doc, ["Registration opens", "Term starts", "Add/drop deadline", "Census date", "Withdrawal deadline", "Term ends", "Grades due"])
	add_heading(doc, "12.3 Course-section validation", 2)
	add_bullets(doc, [
		"Course-based sections require CRN, section number, campus, and delivery method.", "Section and waitlist capacities cannot be negative.",
		"Start time must be earlier than end time.", "Prepared sections use Student Group as the Education-native section record and Course Schedule for dated meetings.",
	])
	add_heading(doc, "12.4 Prepared term snapshot", 2)
	add_table(doc, ["Milestone", "Fall 2026"], [["Registration opens", "2026-04-01"], ["Add/drop deadline", "2026-09-01"], ["Census date", "2026-09-04"], ["Withdrawal deadline", "2026-11-06"], ["Grades due", "2026-12-23"]], [3600, 5760])

	add_heading(doc, "13. Assessment, grading, and report cards", 1)
	add_table(doc, ["Grade", "Minimum score", "Displayed score range"], [["A", "90%", "90% - 100%"], ["B", "80%", "80% - 89.99%"], ["C", "70%", "70% - 79.99%"], ["D", "60%", "60% - 69.99%"], ["F", "0%", "0% - 59.99%"]], [1800, 2700, 4860])
	add_p(doc, "The numeric minimum remains in Education's Percent field because automatic grading maps thresholds to grade codes. The adjacent Score Range provides client-friendly interpretation without changing calculation semantics.")
	add_heading(doc, "13.1 Seeded assessment", 2)
	add_bullets(doc, [
		"Fall 2026 Report Card parent group and Final Examination child group.", "Final Examination criterion and CS 101 assessment plan.",
		"A submitted synthetic Assessment Result for Avery Johnson with a score of 94.", "Report-card tool filters to valid non-root assessment groups.",
	])
	add_heading(doc, "13.2 Compatibility adapter", 2)
	add_bullets(doc, [
		"Overrides Education's preview_report_card endpoint through a supported hook.", "Aggregates Present/Absent attendance using explicit date-bounded SQL compatible with Frappe v16.",
		"Renders the standard Education template with self-contained CSS and rewrites local asset paths for isolated wkhtmltopdf access.",
		"Returns a PDF response named for the student. This is demo report output, not an official transcript.",
	])

	add_heading(doc, "14. Institutional hierarchy and closure materialization", 1)
	add_steps(doc, [
		"Create effective-dated Institution, Campus, Academic Unit Types, and Academic Units.", "Create a draft Structure Version with a bounded validity interval and change summary.",
		"Add exactly one Unit Placement per included unit; every parent must also be placed.", "Submit only after approval reference, interval, uniqueness, institution, unit coverage, depth, and acyclic-forest checks pass.",
		"Closure service obtains a database advisory lock, locks the submitted structure, fingerprints placements, and creates a Building record.",
		"Service bulk-inserts ancestor/descendant/depth rows in batches of 1,000, verifies scope, self rows, depths, fingerprint, and exact expected set.",
		"Build becomes Ready; a forced rebuild supersedes the previous Ready build. Failures retain sanitized evidence and no partial closure rows.",
	])
	add_table(doc, ["Invariant", "Enforcement"], [
		["No cycle / missing parent", "validate_forest"], ["Maximum depth", "Bounded by MAX_STRUCTURE_DEPTH"],
		["Concurrency", "MariaDB GET_LOCK and row lock"], ["Source immutability", "SHA-256 placement fingerprint"],
		["Atomic rows", "Savepoint and rollback"], ["Query safety", "Only Ready builds; institution must match structure"],
	], [3000, 6360])

	page_break(doc)
	add_heading(doc, "15. Demo dataset and workspaces", 1)
	add_table(doc, ["Seed area", "Prepared content"], [
		["Institution", "One institution, one main campus, two academic units"], ["Calendar", "Academic Year 2026-2027 and Fall 2026 milestones"],
		["Catalog", "Two programs and six courses; BSCS contains four required courses"], ["Sections", "Four CRN-based BSCS sections and four schedules"],
		["People", "Three instructors; synthetic applicant/student records"], ["Enrollment", "Program Enrollments, Course Enrollments, and at least one section registration"],
		["Assessment", "One submitted result and report-card scenario"], ["Navigation", "Eight module icons, module sidebars, workspaces, number cards, default workspace"],
	], [2600, 6760])
	add_callout(doc, "Live verification", "At this snapshot all demo, registrar, and US-academic health checks returned ok=true. Six catalog courses, four BSCS sections, four schedules, and at least one section-linked registration were verified.", LIGHT_BLUE)
	add_heading(doc, "15.1 Idempotent setup", 2)
	add_p(doc, "talisma_sis.demo.setup refuses non-demo sites, checks required apps, normalizes Gender options, removes obsolete fields, creates custom fields/property setters, seeds foundation/calendar/catalog/sections/people/assessment, configures workspaces and branding, commits, and returns record counts. Most records use _ensure or explicit lookups so reruns update/reuse rather than blindly duplicate.")

	add_heading(doc, "16. Configuration reference", 1)
	add_heading(doc, "16.1 Application configuration", 2)
	add_table(doc, ["Setting", "Value / location"], [
		["App metadata", "apps/talisma_sis/pyproject.toml and talisma_sis/hooks.py"], ["Python target", ">=3.14; Ruff py314; tab indentation"],
		["Frappe module", "Talisma SIS"], ["Patch order", "talisma_sis/patches.txt post_model_sync"],
		["Assets", "public/css/talisma_demo.css; public/js/*.js; public/icons/*.svg"], ["Translations", "talisma_sis/translations/en.csv"],
	], [2800, 6560])
	add_heading(doc, "16.2 Container variables", 2)
	add_table(doc, ["Variable", "Purpose / default"], [
		["CUSTOM_IMAGE / CUSTOM_TAG", "Select custom application image; demo uses talisma-sis:demo"], ["PULL_POLICY", "Image pull policy; demo uses never for local image"],
		["RESTART_POLICY", "unless-stopped by default"], ["DB_HOST / DB_PORT / DB_PASSWORD", "MariaDB connection; secret value must come from environment/secret store"],
		["REDIS_CACHE / REDIS_QUEUE", "Redis endpoints written by configurator"], ["SOCKETIO_PORT", "9000"],
		["GUNICORN_THREADS / WORKERS / TIMEOUT", "4 / 2 / 120 defaults"], ["FRAPPE_SITE_NAME_HEADER", "demo.talisma.local in demo override"],
		["CLIENT_MAX_BODY_SIZE", "50m default"], ["PROXY_READ_TIMEOUT", "120 default"],
	], [3300, 6060], font_size=8.8)
	add_callout(doc, "Secrets", "Do not place production passwords in Compose files, documentation, source control, images, screenshots, or command history. The local demo credential must be replaced before any shared or network-accessible deployment.", LIGHT_RED)

	add_heading(doc, "17. Security, privacy, and permissions", 1)
	add_bullets(doc, [
		"Server-side DocPerm, has_permission, query conditions, and relationship scoping are the required enforcement points; hidden controls are not authorization.",
		"Student records, grades, advising notes, aid, conduct, disability, identity evidence, and attachments require separate classifications and least privilege.",
		"Faculty access must derive from active teaching/advising assignments; student access must resolve to the authenticated student.",
		"Integration identities require narrow scopes, idempotency, correlation IDs, retries, and audit evidence. Direct database access is prohibited.",
		"The demo uses broad Administrator access for presentation speed and is not evidence of production FERPA compliance.",
	])
	add_heading(doc, "17.1 Architecture decisions", 2)
	add_table(doc, ["ADR", "Decision", "Status"], [["ADR-001", "Frappe Education dependency strategy", "Proposed"], ["ADR-002", "Identity and party model", "Accepted"], ["ADR-003", "Institutional and academic hierarchy", "Accepted; slices gated individually"]], [1500, 5200, 2660])
	add_callout(doc, "Critical legal gate", "Education is GPL-3.0. Qualified legal review of the commercial distribution model remains required before the dependency decision can be accepted.", LIGHT_RED)

	page_break(doc)
	add_heading(doc, "18. Build, deployment, and site operations", 1)
	add_heading(doc, "18.1 Build the demo image", 2)
	add_code(doc, """docker build --build-arg ERPNEXT_IMAGE=talisma-education-base:v16.0.1 `
  --tag talisma-sis:demo --file images/talisma/Containerfile .""")
	add_heading(doc, "18.2 Start or recreate", 2)
	add_code(doc, """$env:CUSTOM_IMAGE='talisma-sis'
$env:CUSTOM_TAG='demo'
$env:PULL_POLICY='never'
docker compose --project-name talisma-demo -f compose.yaml `
  -f overrides/compose.mariadb.yaml -f overrides/compose.redis.yaml `
  -f apps/talisma_sis/demo/compose.override.yaml up -d""")
	add_heading(doc, "18.3 Apply demo setup and verify", 2)
	add_code(doc, """docker compose --project-name talisma-demo -f compose.yaml `
  -f overrides/compose.mariadb.yaml -f overrides/compose.redis.yaml `
  -f apps/talisma_sis/demo/compose.override.yaml exec -T backend `
  bench --site demo.talisma.local execute talisma_sis.demo.setup

docker compose --project-name talisma-demo -f compose.yaml `
  -f overrides/compose.mariadb.yaml -f overrides/compose.redis.yaml `
  -f apps/talisma_sis/demo/compose.override.yaml exec -T backend `
  bench --site demo.talisma.local execute talisma_sis.demo.healthcheck""")
	add_heading(doc, "18.4 Stop safely", 2)
	add_code(doc, """docker compose --project-name talisma-demo -f compose.yaml `
  -f overrides/compose.mariadb.yaml -f overrides/compose.redis.yaml `
  -f apps/talisma_sis/demo/compose.override.yaml down""")
	add_callout(doc, "Destructive option", "Do not add --volumes unless the isolated demo database and site files are intentionally being discarded.", LIGHT_RED)

	add_heading(doc, "19. Testing, health checks, migrations, and recovery", 1)
	add_heading(doc, "19.1 Automated coverage", 2)
	add_bullets(doc, [
		"Institution, campus, unit type, academic unit, structure, placement, and closure controller tests cover validation, immutability, permissions, constraints, lifecycle, and deletion behavior.",
		"Closure tests cover 19,999 generated rows under five seconds, idempotence, rebuild/supersede, lock contention, failed-build evidence, query behavior, and protected service records.",
		"Demo acceptance helpers prove admissions and registration idempotence; health checks verify metadata, navigation, catalog, sections, schedules, and report-card output.",
		"Education compatibility evaluation covers build, install, migrate, services, assets, route behavior, backup, and schema inventory; restore, uninstall, permissions, upgrade, and load tests remain open.",
	])
	add_heading(doc, "19.2 Migration patches", 2)
	add_table(doc, ["Patch", "Purpose"], [
		["add_talisma_campus_unique_constraint", "Institution + campus code uniqueness"],
		["add_talisma_academic_unit_unique_constraint", "Institution + unit code uniqueness"],
		["add_talisma_structure_constraints", "Structure and placement integrity/indexes"],
		["add_talisma_closure_constraints", "Closure build/row integrity and indexes"],
		["ensure_talisma_structure_supersedes_unique", "One successor per superseded structure"],
	], [4200, 5160])
	add_heading(doc, "19.3 Backup and restore", 2)
	add_bullets(doc, [
		"Create a checkpoint with bench --site demo.talisma.local backup --with-files --compress.",
		"Record database, public-file, private-file, and config paths printed by bench.",
		"Stop scheduler and queue workers before restore; restore database and both file archives; restart workers afterward.",
		"Always specify the talisma-demo project and demo override. Never restore into talisma.local.",
	])

	add_heading(doc, "20. Troubleshooting and support", 1)
	add_table(doc, ["Symptom", "Checks / response"], [
		["Old branding/navigation", "Log out, clear boot cache via new login, hard refresh, verify correct port/site"],
		["CSS change not visible", "Confirm image rebuilt, backend/frontend recreated, asset URL contains expected rule"],
		["Form cannot save", "Review mandatory fields and server validation; inspect browser response and backend logs"],
		["Admission button missing", "Applicant must be Approved, unconverted, and on demo site"],
		["No sections available", "Confirm program curriculum, same term/year, Open status, enabled section, and remaining capacity"],
		["Duplicate registration", "Use existing Course Enrollment; different section of same course is rejected"],
		["Report card fails", "Confirm academic year dates, assessment group/result, Chromium/wkhtmltopdf assets, and adapter hook"],
		["Container unhealthy", "docker compose ps/logs; verify DB/Redis health and configurator completion"],
		["Desktop Icon sync warning", "Known Education/Frappe v16 compatibility finding; capture traceback before production acceptance"],
	], [2600, 6760], font_size=8.8)

	page_break(doc)
	add_heading(doc, "21. Client demonstration runbook", 1)
	add_table(doc, ["Time", "Demonstration", "Message"], [
		["0-2 min", "Bryan University home and modules", "Coherent higher-education workspace on an upgrade-safe platform"],
		["2-5 min", "Applicants and admission conversion", "Connected, idempotent applicant-to-student flow"],
		["5-8 min", "Program Enrollment and Register Courses", "Program/term-aware CRN registration and roster linkage"],
		["8-10 min", "Catalog, term milestones, sections", "US-specific academic terminology and validation"],
		["10-12 min", "Assessment, grading range, report card", "Reusable Education foundation with targeted compatibility adaptation"],
	], [1300, 3600, 4460])
	add_heading(doc, "21.1 Prepared story", 2)
	add_bullets(doc, [
		"Use Avery Johnson for the linked Student, Program Enrollment, CRN, and report-card story.",
		"Use another Approved applicant only when deliberately demonstrating conversion; repeated actions should not create duplicates.",
		"Explain the distinction between Program Course, course section (Student Group), schedule, Course Enrollment, and roster membership.",
		"End with the explicit MVP boundary: no production claim for official transcript, degree audit, aid, billing, regulatory reporting, or FERPA certification.",
	])
	add_heading(doc, "21.2 Pre-demo checklist", 2)
	add_bullets(doc, [
		"Services healthy; demo.healthcheck, registrar.healthcheck, and us_academics.healthcheck return ok=true.",
		"Use a private browser window; verify localhost:8082 and Bryan University default workspace.",
		"Confirm synthetic data and no real client/student information.", "Confirm the grading table shows Minimum Score plus Score Range.",
		"Confirm Admission Intake has no Introduction editor and short forms fill the full workspace height.",
	])

	add_heading(doc, "22. Production-readiness gaps and roadmap", 1)
	add_table(doc, ["Phase", "Primary outcome"], [
		["Gate 0", "Accept dependency, licensing, identity, hierarchy, calendar, and FERPA architecture"],
		["Phase 1", "Shared administration, departments, programs, catalog, student/faculty foundations"],
		["Phase 2", "Audited admissions vertical slice with duplicate prevention"],
		["Phase 3", "Versioned curriculum, requisites, offerings, assignments, schedules"],
		["Phase 4", "Deterministic registration: windows, holds, prerequisites, conflicts, waitlists, overrides, census"],
		["Phase 5", "Attendance, assessment, gradebook, final-grade approval, LMS contract"],
		["Phase 6", "Transcript, degree audit, advising, graduation/conferral"],
		["Phase 7", "Student finance and aid with ERPNext Accounts as ledger"],
		["Phase 8", "Portals, analytics, integrations, accessibility, performance, DR, compliance hardening"],
	], [1600, 7760], font_size=8.8)
	add_callout(doc, "Definition of done", "A module is complete only when ownership, lifecycle, permissions, migrations, tests, reports, operations, audit behavior, and integration contracts are approved and verified.", LIGHT_BLUE)


def add_appendices(doc):
	page_break(doc)
	add_heading(doc, "Appendix A. Service and API reference", 1)
	add_table(doc, ["Callable", "Purpose", "Restriction"], [
		["talisma_sis.demo.setup", "Seed/configure complete demo", "demo.talisma.local only"],
		["talisma_sis.demo.healthcheck", "Verify metadata, navigation, branding, ranges, layouts", "demo site"],
		["talisma_sis.demo.dashboard_summary", "Student/applicant/course/faculty counts", "authenticated demo"],
		["talisma_sis.demo.gender_options_query", "Controlled Gender search options", "query endpoint"],
		["talisma_sis.admissions.admit_applicant", "Idempotent applicant conversion", "write permission + approved state + demo"],
		["talisma_sis.registrar.get_registration_summary", "Allowed and registered course summary", "read permission + demo"],
		["talisma_sis.registrar.get_section_options", "Open program/term section options", "demo"],
		["talisma_sis.registrar.register_sections", "Course Enrollment and roster transaction", "create permission + validation + demo"],
		["talisma_sis.registrar.register_courses", "Course-level idempotent registration", "demo; max five"],
		["talisma_sis.report_card.preview_report_card", "Education-compatible PDF output", "overrides upstream method"],
		["talisma_sis.closure.build_closure", "Materialize submitted hierarchy closure", "service-managed records"],
		["talisma_sis.closure.get_descendants / get_ancestors", "Ready-build hierarchy queries", "validated Ready build"],
	], [3700, 3500, 2160], font_size=8.2)

	add_heading(doc, "Appendix B. Key source-file map", 1)
	add_table(doc, ["File", "Responsibility"], [
		["talisma_sis/hooks.py", "Assets, DocType scripts/events, report-card override, post-request redirect"],
		["talisma_sis/demo.py", "Demo guards, metadata, dataset, workspaces, branding, health checks"],
		["talisma_sis/admissions.py", "Applicant-to-student/program orchestration"],
		["talisma_sis/registrar.py", "Program/section registration, roster, summaries, acceptance checks"],
		["talisma_sis/us_academics.py", "US catalog, term, and section validations"],
		["talisma_sis/report_card.py", "Education report-card compatibility adapter"],
		["talisma_sis/closure.py", "Hierarchy closure materialization and queries"],
		["public/js/*.js", "Demo form interactions and dashboard enhancements"],
		["public/css/talisma_demo.css", "Branded responsive and accessible visual system"],
		["demo/compose.override.yaml", "Loopback port and demo site host header"],
		["demo/apps.json", "Pinned ERPNext and Education commits"],
		["images/talisma/Containerfile", "Copies and installs current app into layered image"],
	], [4200, 5160], font_size=8.6)

	add_heading(doc, "Appendix C. Glossary", 1)
	add_table(doc, ["Term", "Meaning"], [
		["Applicant", "Prospective student represented by Student Applicant"], ["Admission Intake", "Configured admission cycle/program eligibility (Student Admission)"],
		["Student", "Academic identity after admission"], ["Program Enrollment", "Student placement in a program and academic period"],
		["Course Enrollment", "Registration in a course, optionally linked to a section"], ["Student Group", "Education-native roster; used as course section"],
		["CRN", "Course Reference Number identifying a section"], ["Academic Unit", "College, school, department, or other governed academic owner"],
		["Structure Version", "Submitted effective-dated academic hierarchy"], ["Closure", "Materialized ancestor/descendant relationships for fast scoped queries"],
		["Census date", "Institutional enrollment snapshot milestone"], ["Idempotent", "Safe repeated operation that does not duplicate the intended record"],
	], [2600, 6760])

	add_heading(doc, "Appendix D. Working-tree and documentation caveats", 1)
	add_bullets(doc, [
		f"Snapshot branch {BRANCH}, commit {COMMIT}; the working tree contains staged, modified, and untracked demo/education files.",
		"This handbook describes that working tree, not a tagged release. Rebuild it after significant code or architecture changes.",
		"The root README's foundation-stage statement is historically accurate but does not include the later isolated demo increment.",
		"Live sample counts can change as users demonstrate workflows. Health-check invariants are more reliable than raw row counts.",
		"The current app package declares Python >=3.14; container runtime compatibility should be confirmed in the release pipeline.",
	])

	add_heading(doc, "Appendix E. Architecture-document index", 1)
	arch_dir = ROOT / "apps" / "talisma_sis" / "docs" / "architecture"
	rows = [[path.name, path.stem.replace("_", " ").replace("-", " ")] for path in sorted(arch_dir.glob("*.md"))]
	add_table(doc, ["Document", "Topic / identifier"], rows, [5200, 4160], font_size=7.8)

	add_heading(doc, "Appendix F. Complete repository manifest", 1)
	add_p(doc, "Searchable file snapshot excluding .git, the isolated document dependency directory, generated render output, and the handbook itself. Upstream Education assets are included because they are part of this working tree.", italic=True)
	excluded = {".git", ".docx-runtime", "project-handbook-render", "__pycache__"}
	files = []
	for path in ROOT.rglob("*"):
		if not path.is_file() or any(part in excluded for part in path.parts):
			continue
		if path == OUT or path.suffix.lower() in {".pyc"}:
			continue
		files.append(path.relative_to(ROOT).as_posix())
	grouped = {}
	for value in sorted(files):
		grouped.setdefault(value.split("/", 1)[0], []).append(value)
	for group, paths in grouped.items():
		add_heading(doc, f"Manifest: {group}", 2)
		add_code(doc, "\n".join(paths))


def audit_docx(path):
	with ZipFile(path) as zf:
		document_xml = zf.read("word/document.xml")
		styles_xml = zf.read("word/styles.xml")
	assert b"Bryan University SIS" in document_xml
	assert b"Heading1" in styles_xml and b"Heading2" in styles_xml
	assert document_xml.count(b"<w:tbl") > 20
	assert document_xml.count(b"<w:tblGrid") == document_xml.count(b"<w:tbl>")


def build():
	doc = Document()
	doc.core_properties.title = "Bryan University SIS Project Handbook"
	doc.core_properties.subject = "Architecture, configuration, workflows, operations, and complete project reference"
	doc.core_properties.author = "Talisma SIS Project"
	doc.core_properties.keywords = "Bryan University, SIS, Frappe, ERPNext, Education, architecture, workflow"
	doc.core_properties.comments = "Generated from the working tree; credentials intentionally omitted."
	configure_styles(doc)
	add_running_furniture(doc)
	add_cover(doc)
	add_contents(doc)
	add_main_content(doc)
	add_appendices(doc)
	OUT.parent.mkdir(parents=True, exist_ok=True)
	doc.save(OUT)
	audit_docx(OUT)
	print(OUT)


if __name__ == "__main__":
	build()
